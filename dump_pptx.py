import pptx

prs = pptx.Presentation(r'C:\Users\ASUS\Desktop\AutoGov_SIH_2026_Submission_Updated.pptx')
with open(r'c:\Users\ASUS\Desktop\26159\autogov_dump.txt', 'w', encoding='utf-8') as f:
    for s_idx, slide in enumerate(prs.slides):
        f.write(f'\n==================== SLIDE {s_idx+1} ====================\n')
        for sh_idx, shape in enumerate(slide.shapes):
            f.write(f'Shape [{sh_idx}] {shape.name} (type {shape.shape_type}):\n')
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        f.write(f'   P: {p.text}\n')
            elif shape.has_table:
                f.write(f'   TABLE {len(shape.table.rows)}x{len(shape.table.columns)}:\n')
                for r_idx, row in enumerate(shape.table.rows):
                    row_txt = [c.text.replace('\n', ' ').strip() for c in row.cells]
                    f.write(f'     Row {r_idx}: {row_txt}\n')
print('Dump complete')
