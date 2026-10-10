import pptx
from pptx.dml.color import RGBColor
from pptx.util import Pt
import os

TEMPLATE_PATH = r'C:\Users\ASUS\Desktop\AutoGov_SIH_2026_Submission_Updated.pptx'
OUTPUT_PATH = r'C:\Users\ASUS\Desktop\SecureMailScope_SIH_2026_Presentation.pptx'
SLIDE1_IMAGE_PATH = r'c:\Users\ASUS\Desktop\26159\slide1_cyber_lock.png'

prs = pptx.Presentation(TEMPLATE_PATH)

def set_run(run, text, font_name=None, font_size_pt=None, bold=None, color_rgb=None):
    run.text = text
    if font_name is not None:
        run.font.name = font_name
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)
    if bold is not None:
        run.font.bold = bold
    if color_rgb is not None:
        run.font.color.rgb = color_rgb

def set_paragraph_simple(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.text = text

def set_paragraph_two_runs(p, text1, text2, bold1=True, bold2=False, sz1=None, sz2=None, c1=None, c2=None, font_name=None):
    if len(p.runs) >= 2:
        set_run(p.runs[0], text1, font_name=font_name, font_size_pt=sz1, bold=bold1, color_rgb=c1)
        set_run(p.runs[1], text2, font_name=font_name, font_size_pt=sz2, bold=bold2, color_rgb=c2)
        for r in p.runs[2:]:
            r.text = ''
    elif len(p.runs) == 1:
        set_run(p.runs[0], text1, font_name=font_name, font_size_pt=sz1, bold=bold1, color_rgb=c1)
        r2 = p.add_run()
        set_run(r2, text2, font_name=font_name, font_size_pt=sz2, bold=bold2, color_rgb=c2)
    else:
        r1 = p.add_run()
        set_run(r1, text1, font_name=font_name, font_size_pt=sz1, bold=bold1, color_rgb=c1)
        r2 = p.add_run()
        set_run(r2, text2, font_name=font_name, font_size_pt=sz2, bold=bold2, color_rgb=c2)

def set_paragraph_three_runs(p, text1, text2, text3, b1=True, b2=True, b3=False, sz=11.5, c1=RGBColor(0x97,0x38,0x35), c2=RGBColor(0,0,0), c3=RGBColor(0x21,0x25,0x29)):
    runs_needed = [ (text1, b1, sz, c1), (text2, b2, sz, c2), (text3, b3, sz, c3) ]
    while len(p.runs) < 3:
        p.add_run()
    for idx, (t, b, s, c) in enumerate(runs_needed):
        set_run(p.runs[idx], t, font_size_pt=s, bold=b, color_rgb=c)
    for r in p.runs[3:]:
        r.text = ''


# ==========================================
# SLIDE 1: TITLE PAGE
# ==========================================
s1 = prs.slides[0]
s1_shapes = s1.shapes

# Header
set_paragraph_simple(s1_shapes[2].text_frame.paragraphs[0], 'SMART INDIA HACKATHON 2026')

# Details Box (Shape 3, TextBox 4)
tb4 = s1_shapes[3].text_frame
set_paragraph_two_runs(tb4.paragraphs[0], '•  Problem Statement ID – ', 'SIH26159', bold1=True, bold2=False, sz1=16.0, sz2=16.0, c1=RGBColor(0,0,0), c2=RGBColor(0x21,0x25,0x29))
set_paragraph_two_runs(tb4.paragraphs[1], '•  Problem Statement Title- ', 'AI-Assisted Cryptographic Security Posture Assessment for Secure Email Communications', bold1=True, bold2=False, sz1=16.0, sz2=16.0, c1=RGBColor(0,0,0), c2=RGBColor(0x21,0x25,0x29))
set_paragraph_two_runs(tb4.paragraphs[2], '•  Theme- ', 'Blockchain & Cybersecurity', bold1=True, bold2=False, sz1=16.0, sz2=16.0, c1=RGBColor(0,0,0), c2=RGBColor(0x21,0x25,0x29))
set_paragraph_two_runs(tb4.paragraphs[3], '•  PS Category- ', 'Software', bold1=True, bold2=False, sz1=16.0, sz2=16.0, c1=RGBColor(0,0,0), c2=RGBColor(0x21,0x25,0x29))
set_paragraph_two_runs(tb4.paragraphs[4], '•  Team ID- ', 'T049', bold1=True, bold2=False, sz1=16.0, sz2=16.0, c1=RGBColor(0,0,0), c2=RGBColor(0x21,0x25,0x29))
set_paragraph_two_runs(tb4.paragraphs[5], '•  Team Name - ', 'Team Innova8 (SecureMailScope)', bold1=True, bold2=False, sz1=16.0, sz2=16.0, c1=RGBColor(0,0,0), c2=RGBColor(0x21,0x25,0x29))

# Replace Slide 1 Graphic (Picture 5) with slide1_cyber_lock.png
for sh in s1_shapes:
    if sh.name == 'Picture 5':
        blip = sh._element.xpath('.//a:blip')
        rId = blip[0].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed']
        image_part = sh.part.related_part(rId)
        with open(SLIDE1_IMAGE_PATH, 'rb') as img_f:
            image_part._blob = img_f.read()


# ==========================================
# SLIDE 2: ADDRESSING THE ISSUE / SOLUTION OVERVIEW
# ==========================================
s2 = prs.slides[1]
s2_shapes = s2.shapes

set_paragraph_simple(s2_shapes[2].text_frame.paragraphs[0], '2')
set_paragraph_simple(s2_shapes[4].text_frame.paragraphs[0], 'SECUREMAILSCOPE')
set_paragraph_simple(s2_shapes[5].text_frame.paragraphs[0], 'AI-Assisted Cryptographic Security Posture Assessment & Immutable Ledger for Email Communications')
set_paragraph_simple(s2_shapes[6].text_frame.paragraphs[0], 'Our Solution')

# 4 Key Feature Bullets (Shapes 8, 10, 12, 14)
c_red = RGBColor(0x97, 0x38, 0x35)
c_dark = RGBColor(0x21, 0x25, 0x29)

set_paragraph_two_runs(s2_shapes[8].text_frame.paragraphs[0], 
                       'Unified Multi-Vector Cryptographic Audit: ', 
                       'Automated live inspection of SPF, DKIM selectors, DMARC enforcement policies, and TLS 1.3 cipher suites.',
                       bold1=True, bold2=False, sz1=12.0, sz2=11.0, c1=c_red, c2=c_dark, font_name='Calibri')

set_paragraph_two_runs(s2_shapes[10].text_frame.paragraphs[0], 
                       'Gradient-Boosted ML Threat Engine: ', 
                       'Quantized GBDT regression model (R²=0.9934) scoring domain posture from 0–100, mapped to CVSS v4.0 criteria.',
                       bold1=True, bold2=False, sz1=12.0, sz2=11.0, c1=c_red, c2=c_dark, font_name='Calibri')

set_paragraph_two_runs(s2_shapes[12].text_frame.paragraphs[0], 
                       'RFC 2822 Deep Header Forensics: ', 
                       'Hop-by-hop relay latency inspection, sender spoofing detection, and forensic authentication header validation.',
                       bold1=True, bold2=False, sz1=12.0, sz2=11.0, c1=c_red, c2=c_dark, font_name='Calibri')

set_paragraph_two_runs(s2_shapes[14].text_frame.paragraphs[0], 
                       'Tamper-Proof Blockchain Audit Ledger: ', 
                       'SHA-256 hashed compliance proofs recorded to an immutable ledger for ISO 27001, SOC 2, and CERT-In audits.',
                       bold1=True, bold2=False, sz1=12.0, sz2=11.0, c1=c_red, c2=c_dark, font_name='Calibri')

# Comparison Table (Shape 15, Table 17)
t17 = s2_shapes[15].table
table_data = [
    ('Problem today', 'How our solution fixes it'),
    ('Fragmented, manual email record checks & hidden spoofing vulnerabilities', 'Instant unified scanner: audits SPF, DKIM, DMARC & TLS in < 2.3 seconds'),
    ("Passive 'p=none' DMARC configurations masquerading as secure defenses", 'Strict policy evaluation & simulated exploit validation highlighting bypass risks'),
    ('Complex, unstandardized cryptographic risk metrics across organizations', 'Calibrated 0-100 GBDT threat scoring normalized to CVSS v4.0 standards'),
    ('Audits lack proof of temporal compliance and tamper-evident logging', 'Immutable blockchain ledger logs cryptographic hashes & verifiable audit proofs')
]
for r_idx, (col0, col1) in enumerate(table_data):
    set_paragraph_simple(t17.rows[r_idx].cells[0].text_frame.paragraphs[0], col0)
    set_paragraph_simple(t17.rows[r_idx].cells[1].text_frame.paragraphs[0], col1)

# 5 Bottom Cards
cards_s2 = [
    (17, 18, '1. Ingest & Probe', 'Target domain or RFC 2822 raw headers ingested via interactive UI or REST API'),
    (20, 21, '2. DNS & TLS Audit', 'Parallel dnspython queries & live SSL handshakes evaluate SPF, DKIM, DMARC & ciphers'),
    (23, 24, '3. ML Threat Scoring', '180-estimator GBDT model computes weighted risk score (0-100) & letter grade (A+-F)'),
    (26, 27, '4. AI Remediation', 'Gemini LLM synthesizes plain-language hardening directives & executive guidance'),
    (29, 30, '5. Ledger & Report', 'SHA-256 state hash anchored to blockchain ledger; exportable compliance reports')
]
for sh_title_idx, sh_body_idx, title, body in cards_s2:
    set_paragraph_simple(s2_shapes[sh_title_idx].text_frame.paragraphs[0], title)
    set_paragraph_simple(s2_shapes[sh_body_idx].text_frame.paragraphs[0], body)


# ==========================================
# SLIDE 3: TECHNICAL APPROACH
# ==========================================
s3 = prs.slides[2]
s3_shapes = s3.shapes

set_paragraph_simple(s3_shapes[2].text_frame.paragraphs[0], 'TECHNICAL APPROACH')
set_paragraph_simple(s3_shapes[4].text_frame.paragraphs[0], '3')
set_paragraph_simple(s3_shapes[5].text_frame.paragraphs[0], 'Tech Stack -')

# Left column: Tech Stack (Shapes 6-13)
tech_stack_items = [
    (6, 'React 19 / Vite', 'Cyberpunk SOC Dashboard & Real-Time Scanner PWA'),
    (7, 'FastAPI / Python', 'High-Concurrency Async API Orchestration Engine'),
    (8, 'dnspython & SSL', 'Live DNS Record Queries & TLS 1.3 Handshake Inspector'),
    (9, 'GradientBoostingRegressor', '180-Estimator Decision Tree Risk Scorer (R²=0.9934)'),
    (10, 'Google Gemini Flash', 'Generative AI Remediation & Hardening Copilot'),
    (11, 'Web3.py / Solidity', 'SHA-256 Immutable Audit Ledger & Proof Verification'),
    (12, 'RFC 2822 Engine', 'Deep Header Forensics & Hop-by-Hop Relay Analysis'),
    (13, 'HIBP & Threat Intel', 'Active Domain Breach & Credential Exposure Intelligence')
]
for sh_idx, p0, p1 in tech_stack_items:
    tf = s3_shapes[sh_idx].text_frame
    set_paragraph_simple(tf.paragraphs[0], p0)
    if len(tf.paragraphs) > 1:
        set_paragraph_simple(tf.paragraphs[1], p1)
    else:
        p_new = tf.add_paragraph()
        set_paragraph_simple(p_new, p1)

# Center Top: Approach 1 (Shapes 15-19)
set_paragraph_simple(s3_shapes[15].text_frame.paragraphs[0], '1   Approach 1 – Real-Time Cryptographic Discovery & Header Forensics')
app1_cards = [
    (16, 'Multi-Selector Discovery', "Probes standard & vendor-specific DKIM selectors; validates SPF record syntax and '-all' strictness."),
    (17, 'Cryptographic Handshake', 'Performs direct TLS handshakes on MX gateways, checking cipher suites, PFS, and cert expiry.'),
    (18, 'RFC 2822 Parser', 'Unpacks hop-by-hop Received headers, calculating transit delays and identifying relay anomalies.'),
    (19, 'Forensic Verdict', 'Flags spoofed envelopes, SPF softfail mismatches, and unauthenticated third-party senders.')
]
for sh_idx, p0, p1 in app1_cards:
    tf = s3_shapes[sh_idx].text_frame
    set_paragraph_simple(tf.paragraphs[0], p0)
    set_paragraph_simple(tf.paragraphs[1], p1)

# Center Bottom: Approach 2 (Shapes 21-25)
set_paragraph_simple(s3_shapes[21].text_frame.paragraphs[0], '2   Approach 2 – Machine Learning Risk Scoring & Blockchain Anchoring')
app2_cards = [
    (22, 'Feature Vector Fusion', 'Extracts 10 calibrated features spanning protocol presence, policy levels, and exposure count.'),
    (23, 'GBDT Risk Inference', 'Gradient Boosted Trees predict normalized risk (MAE=1.16) mapped to CVSS v4.0 severity.'),
    (24, 'Gemini AI Copilot', 'Translates technical cryptographic deficits into copy-paste DNS TXT fixes and mitigation steps.'),
    (25, 'Tamper-Proof Ledger', 'Generates SHA-256 hash of domain audit state and records timestamped tx proof on-chain.')
]
for sh_idx, p0, p1 in app2_cards:
    tf = s3_shapes[sh_idx].text_frame
    set_paragraph_simple(tf.paragraphs[0], p0)
    set_paragraph_simple(tf.paragraphs[1], p1)

# Right Column: Final Output: Officer Cockpit (Shapes 27-32)
set_paragraph_simple(s3_shapes[27].text_frame.paragraphs[0], 'Final Output: SOC Security Cockpit')
output_cards = [
    (28, 'Executive Posture Gauge (0–100)', 'Real-time cryptographic score with letter rating (A+ to F) and status badges.'),
    (29, 'Interactive 9-Step Scanner', 'Live WebSocket/telemetry visualizer showing real-time protocol verification stages.'),
    (30, 'Hop-by-Hop Header Forensic Map', 'Visualizes email transit route, flagging suspicious intermediate relays and IP mismatches.'),
    (31, 'AI Remediation Playbook', 'Context-aware step-by-step instructions to achieve 100/100 cryptographic defense.'),
    (32, 'Verifiable Audit Certificate', 'Cryptographically signed compliance proof with transaction hash and downloadable PDF.')
]
for sh_idx, p0, p1 in output_cards:
    tf = s3_shapes[sh_idx].text_frame
    set_paragraph_simple(tf.paragraphs[0], p0)
    set_paragraph_simple(tf.paragraphs[1], p1)


# ==========================================
# SLIDE 4: FEASIBILITY AND VIABILITY
# ==========================================
s4 = prs.slides[3]
s4_shapes = s4.shapes

set_paragraph_simple(s4_shapes[2].text_frame.paragraphs[0], 'FEASIBILITY AND VIABILITY')
set_paragraph_simple(s4_shapes[4].text_frame.paragraphs[0], '4')
set_paragraph_simple(s4_shapes[5].text_frame.paragraphs[0], 'An enterprise-grade, lightweight cybersecurity solution engineered for turnkey deployment across government ministries, defense bodies, and enterprise email perimeters.')
set_paragraph_simple(s4_shapes[6].text_frame.paragraphs[0], 'Technically feasible, high-performance architecture operating with sub-2.5s assessment latency, requiring zero MX modifications or agent installations on client mail servers.')

# Technical Feasibility (Shape 8 title, Shape 9 textbox)
set_paragraph_simple(s4_shapes[8].text_frame.paragraphs[0], 'Technical Feasibility')
c_blue = RGBColor(0x23, 0x5F, 0xAF)
tf_tf = s4_shapes[9].text_frame
tf_bullets = [
    ('• Non-Intrusive Agentless Assessment: ', 'Conducts external cryptographic audits via RFC-compliant DNS and TLS handshakes—zero server reconfiguration or downtime required.'),
    ('• Lightweight & Deployable ML: ', 'Quantized GBDT pipeline executes inference in < 15ms per domain on standard CPU hardware, eliminating costly GPU infrastructure.'),
    ('• High-Throughput Enterprise API: ', 'FastAPI asynchronous event loop handles 250+ concurrent domain audits with intelligent DNS resolver caching and request pooling.'),
    ('• Zero-PII Compliance Architecture: ', 'Processes public DNS records and sanitized metadata without storing citizen personal data, fully aligning with DPDP Act 2023.')
]
for p_idx, (t1, t2) in enumerate(tf_bullets):
    set_paragraph_two_runs(tf_tf.paragraphs[p_idx], t1, t2, bold1=True, bold2=False, sz1=11.0, sz2=10.0, c1=c_blue, c2=c_dark, font_name='Calibri')

# Potential Risks & Mitigation Strategies (Shape 11 title, Shapes 12-19)
set_paragraph_simple(s4_shapes[11].text_frame.paragraphs[0], 'Potential Risks & Mitigation Strategies')
risk_pairs = [
    (12, 'Risk: Custom DKIM Selectors', 13, 'Fix: Hybrid selector enumeration combines top mail provider dictionaries with RFC 2822 header extraction.'),
    (14, 'Risk: Rate-Limiting & DNS Jitter', 15, 'Fix: Asynchronous multi-threaded resolver pooling with automatic backoff and redundant public DNS fallbacks.'),
    (16, 'Risk: Complex Enterprise Remediation', 17, 'Fix: Gemini AI translates complex cryptanalytic findings into clear, role-specific DNS snippets for admins.'),
    (18, 'Risk: Gas Costs on Blockchain', 19, 'Fix: Merkle-tree state batching anchors thousands of domain audit hashes into single rollup transactions.')
]
for r_sh, r_text, f_sh, f_text in risk_pairs:
    set_paragraph_simple(s4_shapes[r_sh].text_frame.paragraphs[0], r_text)
    set_paragraph_simple(s4_shapes[f_sh].text_frame.paragraphs[0], f_text)

# Practical Implementation & Viability (Shape 21 title, Ovals & Textboxes)
set_paragraph_simple(s4_shapes[21].text_frame.paragraphs[0], 'Practical Implementation & Viability')
viability_points = [
    (23, 'Frictionless Onboarding: ', 'Zero installation required; organizations assess domain security instantly via web or REST API within seconds.'),
    (25, 'Extreme Cost Efficiency: ', 'Drops organizational security posture auditing costs by >90% compared to legacy manual penetration testing.'),
    (27, 'Turnkey Gov & Enterprise Integration: ', 'Containerized microservice readily deploys in NIC MeghRaj, private cloud, or enterprise SOC SIEM pipelines.'),
    (29, 'Tamper-Proof Legal Admissibility: ', 'Cryptographic blockchain state proofs establish incontrovertible compliance records for CERT-In & ISO 27001 audits.'),
    (31, 'Real-Time Threat Intelligence: ', 'Continuous monitoring detects unauthorized DNS drift, certificate expiration, and spoofing vulnerabilities before attackers strike.')
]
for tb_idx, title, body in viability_points:
    tf = s4_shapes[tb_idx].text_frame
    set_paragraph_simple(tf.paragraphs[0], title)
    set_paragraph_simple(tf.paragraphs[1], body)


# ==========================================
# SLIDE 5: IMPACT AND BENEFITS
# ==========================================
s5 = prs.slides[4]
s5_shapes = s5.shapes

set_paragraph_simple(s5_shapes[2].text_frame.paragraphs[0], 'IMPACT AND BENEFITS')
set_paragraph_simple(s5_shapes[4].text_frame.paragraphs[0], '5')
set_paragraph_simple(s5_shapes[5].text_frame.paragraphs[0], 'POTENTIAL IMPACT FLOW')

# 5 Process Flow Cards (Shapes 6-10)
impact_flow = [
    (6, '1. DOMAIN INGESTION', 'Security analyst or admin inputs domain or uploads raw email headers.'),
    (7, '2. MULTI-VECTOR AUDIT', 'Live scanning probes SPF, DKIM selectors, DMARC policies, and TLS ciphers in < 2.3s.'),
    (8, '3. AI & ML EVALUATION', 'GBDT model computes CVSS-mapped threat score; Gemini generates remediation steps.'),
    (9, '4. LEDGER ANCHORING', 'Audit verdict & SHA-256 fingerprint anchored to immutable blockchain ledger.'),
    (10, '5. DEFENSE HARDENING', 'Instant remediation of vulnerabilities eliminates BEC attacks, spoofing, and data exfiltration.')
]
for sh_idx, p0, p1 in impact_flow:
    tf = s5_shapes[sh_idx].text_frame
    set_paragraph_simple(tf.paragraphs[0], p0)
    set_paragraph_simple(tf.paragraphs[1], p1)

set_paragraph_simple(s5_shapes[11].text_frame.paragraphs[0], 'BENEFITS BY PILLAR')

# 5 Pillars (Shapes 13/14, 16/17, 19/20, 22/23, 25/26)
pillars = [
    (13, 'Target Beneficiaries', 14, 
     ['• Government Ministries, Public Sector Undertakings (PSUs) & Defense Organizations',
      '• Enterprise CISOs, SOC Analysts, System Administrators & IT Compliance Officers']),
    (16, 'National & Social Impact', 17, 
     ['• Neutralizes Business Email Compromise (BEC) and spear-phishing campaigns targeting critical infrastructure',
      '• Safeguards citizen trust in official government communications and public service delivery']),
    (19, 'Economic Impact', 20, 
     ['• Prevents multi-crore wire fraud and financial exfiltration caused by spoofed executive emails',
      '• Cuts continuous external security auditing expenditures from lakhs to near-zero automated SaaS costs']),
    (22, 'Operational Impact', 23, 
     ['• Slashing audit turnaround from days of manual DNS inspection to sub-2.5-second automated triage',
      '• 100% comprehensive posture visibility replaces sporadic, periodic spot-checking']),
    (25, 'Technological & Regulatory Impact', 26, 
     ['• Tamper-proof, cryptographically verifiable audit trails fulfilling CERT-In & ISO 27001 cyber compliance',
      '• Explainable AI eliminates guesswork, empowering IT personnel to remediate vulnerabilities instantly'])
]
for title_sh, title_text, body_sh, body_lines in pillars:
    set_paragraph_simple(s5_shapes[title_sh].text_frame.paragraphs[0], title_text)
    tf = s5_shapes[body_sh].text_frame
    for l_idx, line in enumerate(body_lines):
        if l_idx < len(tf.paragraphs):
            set_paragraph_simple(tf.paragraphs[l_idx], line)
        else:
            p_new = tf.add_paragraph()
            set_paragraph_simple(p_new, line)

# Bottom Vision Banner (Shape 27, Rounded Rectangle 29)
tf_vision = s5_shapes[27].text_frame
set_paragraph_simple(tf_vision.paragraphs[0], 'SECUREMAILSCOPE CYBER DEFENSE VISION')
set_paragraph_simple(tf_vision.paragraphs[1], '"An autonomous, verifiable, and explainable cryptographic shield safeguarding national and enterprise email infrastructure against modern spoofing, phishing, and cryptographic obsolescence."')
# P2 has metrics
p_met = tf_vision.paragraphs[2]
metrics = [
    '✓ Sub-2.5s Scan Time   ',
    '✓ R²=0.9934 ML Accuracy   ',
    '✓ CVSS v4.0 Mapped   ',
    '✓ Immutable Blockchain Proof'
]
while len(p_met.runs) < len(metrics):
    p_met.add_run()
for idx, m_text in enumerate(metrics):
    set_run(p_met.runs[idx], m_text)
for r in p_met.runs[len(metrics):]:
    r.text = ''


# ==========================================
# SLIDE 6: RESEARCH AND REFERENCES
# ==========================================
s6 = prs.slides[5]
s6_shapes = s6.shapes

set_paragraph_simple(s6_shapes[2].text_frame.paragraphs[0], 'RESEARCH  AND REFERENCES')
set_paragraph_simple(s6_shapes[4].text_frame.paragraphs[0], '6')

tb_ref = s6_shapes[5].text_frame

# Section 1: Standards
set_paragraph_simple(tb_ref.paragraphs[0], '1. International Email Cryptography & Security Standards')
set_paragraph_three_runs(tb_ref.paragraphs[1], '•  ', 'IETF RFC 7208 (SPF): ', 'Sender Policy Framework for Authorizing Use of Domains in Email — Formal specifications for DNS-based host identity authorization.')
set_paragraph_three_runs(tb_ref.paragraphs[2], '•  ', 'IETF RFC 6376 (DKIM): ', 'DomainKeys Identified Mail Signatures — Public-key cryptographic authentication for email message integrity.')
set_paragraph_three_runs(tb_ref.paragraphs[3], '•  ', 'IETF RFC 7489 (DMARC): ', 'Domain-based Message Authentication, Reporting, and Conformance — Policy enforcement and reporting framework against domain spoofing.')
set_paragraph_three_runs(tb_ref.paragraphs[4], '•  ', 'IETF RFC 8446 (TLS 1.3): ', 'The Transport Layer Security Protocol Version 1.3 — Modern forward-secret cryptographic handshakes and deprecated cipher protection.')

# Section 2: Policies & Regulatory
set_paragraph_simple(tb_ref.paragraphs[5], '2. Cybersecurity Guidelines & Regulatory Directives')
set_paragraph_three_runs(tb_ref.paragraphs[6], '•  ', 'CERT-In Cybersecurity Directions (No. 20(3)/2022-CERT-In): ', 'Mandatory logging, security incident reporting, and strict cryptographic posture requirements.')
set_paragraph_three_runs(tb_ref.paragraphs[7], '•  ', 'Digital Personal Data Protection (DPDP) Act 2023: ', 'Mandatory data minimization, purpose limitation, and zero-PII architectural standards for citizen document handling.')
set_paragraph_three_runs(tb_ref.paragraphs[8], '•  ', 'ISO/IEC 27001:2022 (Annex A.8.20 & A.8.24): ', 'Network security, information transfer, and use of cryptography controls for organizational email systems.')
set_paragraph_three_runs(tb_ref.paragraphs[9], '•  ', 'NIST Special Publication 800-177 (Rev. 1): ', 'Trustworthy Email — Federal recommendations on cryptographic mechanisms for securing email infrastructures.')

# Section 3: ML & Ledgers
set_paragraph_simple(tb_ref.paragraphs[10], '3. Machine Learning, Threat Scoring & Cryptographic Ledgers')
set_paragraph_three_runs(tb_ref.paragraphs[11], '•  ', 'FIRST CVSS v4.0 Specification: ', 'Common Vulnerability Scoring System v4.0: Modern metric guidelines for quantitative cryptographic vulnerability assessment.')
set_paragraph_three_runs(tb_ref.paragraphs[12], '•  ', 'Friedman, J. H. (2001): ', 'Greedy Function Approximation: A Gradient Boosting Machine (Annals of Statistics) — Theoretical foundation for non-linear risk regressors.')
set_paragraph_three_runs(tb_ref.paragraphs[13], '•  ', 'Lundberg & Lee (2017): ', 'A Unified Approach to Interpreting Model Predictions (SHAP) — Explaining cryptographic risk feature contributions for actionable remediation.')

# Save final presentation
prs.save(OUTPUT_PATH)
print(f'Final presentation created successfully at: {OUTPUT_PATH}')
