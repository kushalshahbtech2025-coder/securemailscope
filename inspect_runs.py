import pptx

prs = pptx.Presentation(r'C:\Users\ASUS\Desktop\AutoGov_SIH_2026_Submission_Updated.pptx')
with open(r'c:\Users\ASUS\Desktop\26159\runs_formatting.txt', 'w', encoding='utf-8') as f:
    for s_idx, slide in enumerate(prs.slides):
        f.write(f'\n=== SLIDE {s_idx+1} ===\n')
        for sh_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                tf = shape.text_frame
                for p_idx, p in enumerate(tf.paragraphs):
                    if p.text.strip():
                        runs_info = []
                        for r in p.runs:
                            font_name = r.font.name
                            font_size = r.font.size.pt if r.font.size else None
                            bold = r.font.bold
                            color = r.font.color.rgb if (r.font.color and r.font.color.type == 1) else None
                            runs_info.append(f'"{r.text}" ({font_name}, {font_size}pt, b={bold}, c={color})')
                        f.write(f'  Shape [{sh_idx}] {shape.name} P{p_idx}: ' + ' + '.join(runs_info) + '\n')
            elif shape.has_table:
                f.write(f'  Shape [{sh_idx}] {shape.name} TABLE:\n')
                for r_idx, row in enumerate(shape.table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        for p in cell.text_frame.paragraphs:
                            if p.text.strip():
                                f.write(f'    Cell ({r_idx},{c_idx}): "{p.text}"\n')
print('Formatting inspection saved to runs_formatting.txt')
